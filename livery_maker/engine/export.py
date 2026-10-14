"""Build every aircraft in a project and write the fleet spreadsheet."""
from __future__ import annotations

import os
import zipfile
from typing import Callable, Dict, List, Optional

from .common import FleetEntry, Project, noop
from .templates import TEMPLATES


def validate(project: Project) -> List[str]:
    problems = []
    if not project.output_dir:
        problems.append('Choose an output folder.')
    regs = [f.reg.strip().upper() for f in project.fleet]
    for r in set(regs):
        if r and regs.count(r) > 1:
            problems.append(f'Registration {r} is used more than once.')
    for f in project.fleet:
        if not f.reg.strip():
            problems.append(f'An aircraft ({TEMPLATES[f.template].name}) has no registration.')
    for tid in sorted({f.template for f in project.fleet}):
        for p in TEMPLATES[tid].check_source(project.sources.get(tid, '')):
            problems.append(f'{TEMPLATES[tid].name}: {p}')
    return problems


def warnings(project: Project) -> List[str]:
    out = []
    for tid in sorted({f.template for f in project.fleet}):
        out += [f'{TEMPLATES[tid].name}: {w}' for w in TEMPLATES[tid].source_warnings(project.sources.get(tid, ''))]
    return out


def build_one(project: Project, entry: FleetEntry, progress=noop):
    t = TEMPLATES[entry.template]
    return t.build(project.sources[entry.template], project.design, entry, project.output_dir, progress)


def build_all(project: Project, progress: Callable[[str], None] = noop,
              on_done: Optional[Callable[[FleetEntry, Dict], None]] = None, cancelled=lambda: False,
              make_zips: bool = False) -> List[str]:
    os.makedirs(project.output_dir, exist_ok=True)
    folders = []
    order = sorted(project.fleet, key=lambda f: list(TEMPLATES).index(f.template))   # same kit back to back (cache)
    for i, entry in enumerate(order, 1):
        if cancelled():
            break
        progress(f'[{i}/{len(order)}] {entry.reg}')
        prev = build_one(project, entry, progress)
        folder = os.path.join(project.output_dir, TEMPLATES[entry.template].folder_name(project.design, entry))
        folders.append(folder)
        if make_zips:
            zip_folder(folder)
        if on_done:
            on_done(entry, prev)
    write_fleet_xlsx(project, os.path.join(project.output_dir, f'{project.design.airline} fleet.xlsx'))
    return folders


def zip_folder(folder: str) -> str:
    path = folder.rstrip('/') + '.zip'
    base = os.path.dirname(folder)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for dp, _dn, fns in os.walk(folder):
            for fn in fns:
                full = os.path.join(dp, fn)
                z.write(full, os.path.relpath(full, base))
    return path


FILLS = {'A319': '9DC3E6', 'A321': 'D5A6E6', 'A21N': 'C6E0B4', 'E145': 'F8CBAD', 'Q400': 'FFFF00', 'E190': 'F4B183'}


def write_fleet_xlsx(project: Project, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook(); ws = wb.active; ws.title = 'Fleet'
    hdr = ['Reg', 'Sim', 'Airline', 'Make', 'Model', 'CEO/NEO', 'Engines', 'SELCAL', 'Base', 'Livery folder']
    ws.append(hdr)
    for f in sorted(project.fleet, key=lambda e: e.reg):
        t = TEMPLATES[f.template]
        v = t.variant_key(f)
        model = t.model_for(f)
        gen = v.get('gen', '–')
        eng = f.variant.split(' ', 1)[-1] if t.id == 'toliss_a321' else f.variant
        ws.append([f.reg, t.sim, project.design.airline, t.make, model, gen, eng, f.selcal, project.design.home_base,
                   t.folder_name(project.design, f)])
    font = Font(name='Arial', size=12); bold = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    thin = Side(style='thin', color='70AD47')
    for c in ws[1]:
        c.font = bold; c.fill = PatternFill('solid', fgColor='00B050')
    for row in ws.iter_rows(min_row=2):
        fill = FILLS.get(row[4].value, 'FFFFFF')
        for c in row:
            c.font = font; c.fill = PatternFill('solid', fgColor=fill); c.border = Border(top=thin, bottom=thin)
    for i, w in enumerate([12, 10, 14, 11, 9, 11, 13, 10, 9, 40]):
        ws.column_dimensions[chr(65 + i)].width = w
    if ws.max_row > 1:
        tbl = Table(displayName='Fleet', ref=f'A1:{chr(64 + len(hdr))}{ws.max_row}')
        tbl.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=False)
        ws.add_table(tbl)
    ws.freeze_panes = 'A2'
    wb.save(path)
